"""Excel exports (openpyxl): leaderboard, master recruitment tracker, and
hiring-manager candidate summaries — all fields exactly per the brief."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import config
from ..models import Candidate, Evaluation, Job

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _autosize(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _fmt_dt(value) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else ""


def _latest_eval(candidate: Candidate) -> Evaluation | None:
    return candidate.evaluations[0] if candidate.evaluations else None


def export_leaderboard(job: Job, rows: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leaderboard"
    headers = [
        "Rank", "Candidate ID", "Candidate", "Overall Match %", "Mandatory Criteria",
        "Relevant Experience (yrs)", "Key Strengths", "Gaps", "Risk Flags",
        "AI Recommendation", "Confidence", "Evidence", "Recruiter Decision", "Status",
    ]
    ws.append(headers)
    for row in rows:
        evidence = "\n".join(
            f"[{e['status']}] {e['criterion']}: \"{e['evidence'][:150]}\""
            for e in row.get("evidence", [])
        )
        ws.append([
            row["rank"], row["candidate_code"], row["candidate"], row["overall_match"],
            row["mandatory_criteria"], row["relevant_experience_years"],
            "; ".join(row.get("key_strengths", [])), "; ".join(row.get("gaps", [])),
            "; ".join(row.get("risk_flags", [])), row["recommendation"],
            row["confidence"], evidence, row.get("recruiter_decision", ""),
            row.get("status", ""),
        ])
    _style_header(ws, len(headers))
    _autosize(ws, [6, 12, 24, 12, 16, 12, 34, 34, 28, 18, 11, 60, 16, 16])
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = WRAP

    path = config.EXPORT_DIR / f"leaderboard_{job.job_code}_{_stamp()}.xlsx"
    wb.save(path)
    return path


TRACKER_HEADERS = [
    "Job ID", "Job Title", "Candidate ID", "Candidate Name", "Resume Source",
    "Date Received", "Date Screened", "AI Score", "AI Recommendation",
    "Recruiter Decision", "Screening Status", "Interview Stage", "Interview Date",
    "Feedback Status", "Offer Status", "Joining Status", "Rejection Reason",
    "Next Action", "Owner", "Last Updated",
]


def export_tracker(jobs: list[Job]) -> Path:
    """Master recruitment tracker across jobs, one row per candidate."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Recruitment Tracker"
    ws.append(TRACKER_HEADERS)
    for job in jobs:
        for candidate in job.candidates:
            evaluation = _latest_eval(candidate)
            ws.append([
                job.job_code, job.title, candidate.candidate_code, candidate.full_name,
                candidate.source, _fmt_dt(candidate.received_at),
                _fmt_dt(candidate.processed_at),
                evaluation.overall_score if evaluation else "",
                evaluation.recommendation if evaluation else "",
                candidate.recruiter_decision, candidate.status,
                candidate.interview_stage, _fmt_dt(candidate.interview_date),
                candidate.feedback_status, candidate.offer_status,
                candidate.joining_status, candidate.rejection_reason,
                candidate.next_action, candidate.owner, _fmt_dt(candidate.updated_at),
            ])
    _style_header(ws, len(TRACKER_HEADERS))
    _autosize(ws, [10, 26, 12, 22, 16, 17, 17, 9, 18, 16, 18, 16, 17, 14, 12, 13, 30, 24, 14, 17])
    path = config.EXPORT_DIR / f"recruitment_tracker_{_stamp()}.xlsx"
    wb.save(path)
    return path


SUMMARY_HEADERS = [
    "Candidate", "Candidate ID", "Job", "Client", "Current Role & Employer",
    "Total Experience", "Relevant Experience", "Core Skills",
    "Industry/Client Exposure", "Key Achievements", "Current Compensation",
    "Expected Compensation", "Notice Period", "Location Confirmed",
    "Shift Confirmed", "Recruiter Observations", "Areas to Probe",
    "Reason for Recommendation", "AI Score", "AI Recommendation",
]


def export_hm_summaries(job: Job, summaries: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Summaries"
    ws.append(SUMMARY_HEADERS)
    for s in summaries:
        confirm = {True: "Yes", False: "No", None: "TBC"}
        ws.append([
            s.get("candidate_name"), s.get("candidate_code"), s.get("job_title"),
            s.get("client"), s.get("current_role_and_employer"),
            s.get("total_experience"), s.get("relevant_experience"),
            ", ".join(s.get("core_skills", [])), s.get("industry_client_exposure"),
            "\n".join(f"• {a}" for a in s.get("key_achievements", [])),
            s.get("current_compensation"), s.get("expected_compensation"),
            s.get("notice_period"), confirm[s.get("location_confirmed")],
            confirm[s.get("shift_confirmed")], s.get("recruiter_observations"),
            "\n".join(f"• {p}" for p in s.get("areas_to_probe", [])),
            s.get("reason_for_recommendation"), s.get("ai_score"),
            s.get("ai_recommendation"),
        ])
    _style_header(ws, len(SUMMARY_HEADERS))
    _autosize(ws, [22, 12, 24, 14, 30, 12, 12, 34, 24, 44, 14, 14, 12, 10, 10, 40, 40, 44, 8, 16])
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = WRAP
    path = config.EXPORT_DIR / f"hm_summaries_{job.job_code}_{_stamp()}.xlsx"
    wb.save(path)
    return path


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
